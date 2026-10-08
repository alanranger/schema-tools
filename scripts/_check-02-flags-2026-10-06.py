from pathlib import Path
import openpyxl

p = Path(r'G:\Dropbox\alan ranger photography\Website Code\alan-shared-resources\csv processed') / '02 – products_cleaned.xlsx'
wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print('headers:', headers)
idx = {h: i for i, h in enumerate(headers) if h}
print('flag cols:', [h for h in headers if h and any(k in str(h).lower() for k in ('exclude', 'manual', 'manifest', 'paste'))])
# find slug col
slug_keys = [h for h in headers if h and 'slug' in str(h).lower()]
url_keys = [h for h in headers if h and 'url' in str(h).lower()]
print('slug_keys', slug_keys, 'url_keys', url_keys)
si = idx.get(slug_keys[0]) if slug_keys else 0
for row in ws.iter_rows(min_row=2, values_only=True):
    slug = str(row[si] or '')
    if any(k in slug.lower() for k in ('mentor', '2hr', '2-hour', 'warwickshire-woodland', 'four-private', 'face-to-face')):
        data = {h: row[idx[h]] for h in headers if h in idx}
        interesting = {k: v for k, v in data.items() if k and (k in ('product_slug','slug','url','name','product_name') or 'exclude' in k.lower() or 'manual' in k.lower())}
        print(interesting)
wb.close()

# compile check generate
import py_compile
py_compile.compile(r'G:\Dropbox\alan ranger photography\Website Code\Schema Tools\scripts\generate-product-schema.py', doraise=True)
print('generate-product-schema.py compiles OK')
