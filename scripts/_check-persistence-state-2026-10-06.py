from pathlib import Path
import re
import csv
from collections import Counter

mg = Path(r'G:\Dropbox\alan ranger photography\Website Code\Schema Tools\scripts\match-google-reviews.py')
gen = Path(r'G:\Dropbox\alan ranger photography\Website Code\Schema Tools\scripts\generate-product-schema.py')
dep = Path(r'G:\Dropbox\alan ranger photography\Website Code\Schema Tools\scripts\deploy-product-schemas-cli.mjs')
copy = Path(r'G:\Dropbox\alan ranger photography\Website Code\Schema Tools\scripts\_copy-schemas-no-git-2026-10-06.mjs')

for label, p, pats in [
    ('match', mg, [r'15-review', r'override', r'apply_attribution']),
    ('gen', gen, [r'manifest-policy', r'exclude_from_manifest', r'manual_paste', r'write_manifest']),
    ('deploy', dep, [r'manifest-policy']),
    ('copy', copy, [r'manifest-policy']),
]:
    text = p.read_text(encoding='utf-8', errors='replace')
    print(f'=== {label} {p.name} ({len(text)} chars) ===')
    for pat in pats:
        hits = [(i + 1, line.strip()[:140]) for i, line in enumerate(text.splitlines()) if re.search(pat, line, re.I)]
        print(f'  {pat}: {len(hits)} hits')
        for h in hits[:8]:
            print(f'    {h[0]}: {h[1]}')

ov = Path(r'G:\Dropbox\alan ranger photography\Website Code\alan-shared-resources\csv processed\15-review-attribution-overrides.csv')
with ov.open(encoding='utf-8') as f:
    rows = list(csv.DictReader(f))
print(f'overrides rows: {len(rows)}')
print('scopes:', Counter(x['scope'] for x in rows))
print('barbara:', [x for x in rows if 'Barbara' in x.get('reviewer_name', '')])
print('business count:', sum(1 for x in rows if x.get('scope') == 'business_level'))
print('f2f/four:', [x.get('reviewer_name') + '->' + (x.get('product_slug') or '') for x in rows if 'four-private' in (x.get('product_slug') or '') or 'face-to-face' in (x.get('product_slug') or '')])

# 02 flags via openpyxl
import openpyxl
wb = openpyxl.load_workbook(r'G:\Dropbox\alan ranger photography\Website Code\alan-shared-resources\csv processed\02 - products_cleaned.xlsx', read_only=True)
ws = wb.active
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print('02 headers has flags:', [h for h in headers if h and ('exclude' in str(h).lower() or 'manual' in str(h).lower() or 'manifest' in str(h).lower())])
idx = {h: i for i, h in enumerate(headers)}
for row in ws.iter_rows(min_row=2, values_only=True):
    slug = row[idx.get('product_slug') or idx.get('slug') or 0]
    if not slug:
        continue
    s = str(slug)
    if 'mentoring' in s or '2hr' in s or '2-hour' in s or 'two-hour' in s or 'warwickshire-woodland' in s:
        extras = {}
        for k in ('exclude_from_manifest', 'manual_paste_page', 'product_slug', 'slug', 'name'):
            if k in idx:
                extras[k] = row[idx[k]]
        if extras:
            print('02 row:', extras)
wb.close()
